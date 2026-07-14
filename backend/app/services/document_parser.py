from dataclasses import dataclass
from pathlib import Path
import pymupdf
import pytesseract
from docx import Document
from PIL import Image
from app.core.config import settings

@dataclass(frozen=True)
class ParsedSection:
    text: str
    page_number: int | None = None
    used_ocr: bool = False

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

def inspect_document(path: Path) -> int:
    if path.suffix.lower() == ".pdf":
        with pymupdf.open(path) as document:
            return len(document)
    return 1

def parse_document(path: Path, progress_callback=None) -> list[ParsedSection]:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return _parse_pdf(path, progress_callback)
    if extension == ".docx":
        return _parse_docx(path)
    if extension in {".txt", ".md"}:
        return _parse_text(path)
    raise ValueError(f"Unsupported file extension: {extension}")

def _parse_pdf(path: Path, progress_callback=None) -> list[ParsedSection]:
    pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
    sections = []

    with pymupdf.open(path) as document:
        if len(document) > settings.knowledge_max_pages:
            raise ValueError(
                f"PDF has {len(document)} pages; maximum is {settings.knowledge_max_pages}."
            )

        for index, page in enumerate(document):
            text = page.get_text("text", sort=True).strip()
            used_ocr = False

            if len(text) < settings.knowledge_min_text_chars_per_page:
                text = _ocr_page(page)
                used_ocr = True

            if text.strip():
                sections.append(
                    ParsedSection(text=text.strip(), page_number=index + 1, used_ocr=used_ocr)
                )

            if progress_callback:
                progress_callback(index + 1, len(document), used_ocr)

    return sections

def _ocr_page(page) -> str:
    scale = settings.ocr_dpi / 72
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)

    try:
        return pytesseract.image_to_string(image, lang=settings.ocr_languages).strip()
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError(
            "OCR is required, but Tesseract was not found. Set TESSERACT_CMD in .env."
        ) from exc
    except pytesseract.TesseractError as exc:
        raise RuntimeError(f"Tesseract OCR failed. Check OCR_LANGUAGES: {exc}") from exc

def _parse_docx(path: Path) -> list[ParsedSection]:
    document = Document(path)
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    text = "\n".join(blocks).strip()
    return [ParsedSection(text=text)] if text else []

def _parse_text(path: Path) -> list[ParsedSection]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [ParsedSection(text=text)] if text else []
